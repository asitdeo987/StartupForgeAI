from io import BytesIO

from pptx import Presentation
from pptx.util import Inches, Pt


def generate_ppt(pitch: dict):

    prs = Presentation()

    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    #
    # Cover Slide
    #

    slide_layout = prs.slide_layouts[0]

    slide = prs.slides.add_slide(slide_layout)

    slide.shapes.title.text = pitch["startup_name"]

    slide.placeholders[1].text = pitch["tagline"]

    #
    # Remaining Slides
    #

    for item in pitch["slides"]:

        layout = prs.slide_layouts[1]

        slide = prs.slides.add_slide(layout)

        slide.shapes.title.text = item["title"]

        body = slide.placeholders[1].text_frame

        body.clear()

        for point in item["content"]:

            p = body.add_paragraph()

            p.text = point

            p.level = 0

            p.font.size = Pt(22)

    buffer = BytesIO()

    prs.save(buffer)

    buffer.seek(0)

    return buffer